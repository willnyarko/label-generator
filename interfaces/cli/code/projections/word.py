"""Write the label sheets as Word files, laid out to the sheets in ``config/label_sheets.toml``."""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import _Cell

from lrr.box_kinds import STOCKS, group_boxes_by_kind
from models.lrr import BoxLRR, FolderLRR, LrrBundle
from projections.label_content import ancestor_chain, box_folder_line, box_heading, box_range_text
from projections.label_sheets import FOLDER_SHEET, SHEET_SPECS
from projections.options import DEFAULT_RENDER_OPTIONS, BoxKind, RenderOptions
from projections.typography import typography


def write_word_outputs(
    bundle: LrrBundle,
    output_dir: Path,
    stem: str,
    *,
    options: Optional[RenderOptions] = None,
) -> list[Path]:
    opts = options or DEFAULT_RENDER_OPTIONS
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []

    if bundle.folders:
        folder_suffix = "_left_labels" if opts.folder_layout == "left" else "_folder_labels"
        folder_path = output_dir / f"{stem}{folder_suffix}.docx"
        render_folder_docx(bundle, layout=opts.folder_layout).save(str(folder_path))
        paths.append(folder_path)

    groups = group_boxes_by_kind(bundle.boxes, choice=opts.box_kind)
    for kind, boxes in groups.items():
        stock = STOCKS[kind]
        box_path = output_dir / f"{stem}{stock.suffix}.docx"
        render_box_docx(boxes, kind=kind).save(str(box_path))
        paths.append(box_path)

    return paths


def render_folder_docx(bundle: LrrBundle, *, layout: str = "regular") -> Document:
    spec = SHEET_SPECS[FOLDER_SHEET]
    doc = Document()
    _setup_page(doc, spec)
    if layout == "left":
        _paginate_folder_left(doc, spec, list(bundle.folders))
    else:
        _paginate_folder_regular(doc, spec, list(bundle.folders))
    return doc


def render_box_docx(boxes: list[BoxLRR], *, kind: BoxKind = "full") -> Document:
    stock = STOCKS[kind]
    spec = SHEET_SPECS[stock.sheet]
    doc = Document()
    _setup_page(doc, spec)
    _paginate_boxes(doc, spec, boxes, vertical=stock.vertical, type_=typography()["box"][kind])
    return doc


def _setup_page(doc: Document, spec: dict) -> None:
    section = doc.sections[0]
    section.page_width = Inches(spec["page_width_in"])
    section.page_height = Inches(spec["page_height_in"])
    section.left_margin = Inches(spec["margin_side_in"])
    section.right_margin = Inches(spec["margin_side_in"])
    section.top_margin = Inches(spec["margin_top_in"])
    section.bottom_margin = Inches(0)


def _configure_table(table, spec: dict) -> None:
    table.allow_autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w = spec["label_width_in"]
    g = spec["gutter_in"]
    h = spec["label_height_in"]
    col_widths = [w, g, w]
    for row in table.rows:
        row.height = Inches(h)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for i, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[i])


def _thin_paragraph(doc: Document):
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(1)
    return para


def _sheets(doc: Document, spec: dict, items: list, per_page: int):
    chunks = [items[i : i + per_page] for i in range(0, len(items), per_page)] or [[]]
    for index, chunk in enumerate(chunks):
        if index:
            _thin_paragraph(doc).add_run().add_break(WD_BREAK.PAGE)
        table = doc.add_table(rows=spec["rows"], cols=3)
        _configure_table(table, spec)
        yield table, chunk
    _thin_paragraph(doc)


def _slot(spec: dict, i: int) -> tuple[int, int]:
    r, slot = divmod(i, spec["cols"])
    return r, 0 if slot == 0 else 2


def _paginate_folder_regular(doc: Document, spec: dict, items: list[FolderLRR]) -> None:
    for table, chunk in _sheets(doc, spec, items, spec["rows"]):
        for i, item in enumerate(chunk):
            _fill_folder_left_cell(table.cell(i, 0), item)
            _fill_folder_right_cell(table.cell(i, 2), item)


def _paginate_folder_left(doc: Document, spec: dict, items: list[FolderLRR]) -> None:
    for table, chunk in _sheets(doc, spec, items, spec["cols"] * spec["rows"]):
        for i, item in enumerate(chunk):
            _fill_folder_left_cell(table.cell(*_slot(spec, i)), item)


def _paginate_boxes(doc: Document, spec: dict, items: list[BoxLRR], *, vertical: bool, type_: dict) -> None:
    for table, chunk in _sheets(doc, spec, items, spec["cols"] * spec["rows"]):
        for i, item in enumerate(chunk):
            cell = table.cell(*_slot(spec, i))
            if vertical:
                _set_vertical_text(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _fill_box_cell(cell, item, type_)


def _set_vertical_text(cell: _Cell) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    text_dir = OxmlElement("w:textDirection")
    text_dir.set(qn("w:val"), "btLr")
    tcPr.append(text_dir)


def _line(
    cell: _Cell,
    text: str,
    *,
    type_: dict,
    center: bool = False,
    first: bool = False,
) -> None:
    para = cell.paragraphs[0] if first else cell.add_paragraph()
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text or "")
    run.font.size = Pt(type_["size"])
    run.font.name = type_["font"]


def _fill_folder_left_cell(cell: _Cell, row: FolderLRR) -> None:
    t = typography()["folder"]
    _line(cell, str(row.call_number), type_=t, first=True)
    _line(cell, box_folder_line(row), type_=t)
    _line(cell, str(row.collection), type_=t)
    if row.series:
        _line(cell, row.series, type_=t)


def _fill_folder_right_cell(cell: _Cell, row: FolderLRR) -> None:
    t = typography()["folder"]
    _line(cell, ancestor_chain(row), type_=t, first=True)
    _line(cell, str(row.folder_title), type_=t)
    if row.folder_dates:
        _line(cell, str(row.folder_dates), type_=t)


def _fill_box_cell(cell: _Cell, row: BoxLRR, t: dict) -> None:
    _line(cell, str(row.repository), type_=t, center=True, first=True)
    _line(cell, str(row.collection), type_=t, center=True)
    _line(cell, str(row.call_number), type_=t, center=True)
    if row.series:
        _line(cell, row.series, type_=t, center=True)
    _line(cell, box_heading(row), type_=t, center=True)
    folders = box_range_text(row)
    if folders:
        _line(cell, folders, type_=t, center=True)
